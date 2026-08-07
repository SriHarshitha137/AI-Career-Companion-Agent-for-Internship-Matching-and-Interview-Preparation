"""Extract readable text from uploaded PDF and DOCX files."""

from __future__ import annotations

from io import BytesIO

import pdfplumber
from docx import Document


class TextExtractionError(Exception):
    """Raised when a resume cannot be read as its advertised format."""


def extract_text(file_bytes: bytes, extension: str) -> str:
    """Return resume text, preserving page/paragraph/table boundaries where possible."""
    if not file_bytes:
        raise TextExtractionError("The uploaded file is empty.")

    try:
        if extension.lower() == ".pdf":
            # pdfplumber keeps each page separate, useful when the resume has columns.
            with pdfplumber.open(BytesIO(file_bytes)) as pdf:
                pages = [page.extract_text(layout=True) or "" for page in pdf.pages]
            text = "\n\n--- Page Break ---\n\n".join(pages)
        elif extension.lower() == ".docx":
            document = Document(BytesIO(file_bytes))
            paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
            # Tables commonly contain education, skills, and contact information.
            table_rows = []
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip().replace("\n", " | ") for cell in row.cells]
                    if any(cells):
                        table_rows.append(" | ".join(cells))
            text = "\n".join(paragraphs + table_rows)
        else:
            raise TextExtractionError("Only .pdf and .docx resume files are supported.")
    except TextExtractionError:
        raise
    except Exception as exc:
        raise TextExtractionError(
            "Could not read this file. It may be corrupted or not a valid resume document."
        ) from exc

    if not text.strip():
        raise TextExtractionError(
            "No readable text was found in the resume. For a scanned PDF, run OCR first."
        )
    return text.strip()
