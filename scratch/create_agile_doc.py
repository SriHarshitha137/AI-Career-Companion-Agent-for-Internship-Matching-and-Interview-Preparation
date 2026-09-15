"""Generate professional Agile_Project_Documentation.docx matching the template."""

from __future__ import annotations

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor


def set_cell_background(cell, hex_color: str) -> None:
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def style_table(table, col_widths, headers, rows):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1F2937")  # Dark Slate
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.bold = True
            run.font.name = "Calibri"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(255, 255, 255)

    # Data Rows
    for r_idx, row_data in enumerate(rows):
        row_cells = table.add_row().cells
        bg_color = "F9FAFB" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=90, bottom=90, left=140, right=140)
            p = row_cells[c_idx].paragraphs[0]
            for run in p.runs:
                run.font.name = "Calibri"
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(31, 41, 55)

    # Widths
    for row in table.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = Inches(w)


def build_document():
    doc = Document()

    # Page Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Page 1: Title Page
    p_top = doc.add_paragraph()
    p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_top.add_run("Agile Project Documentation\n\n")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(14)
    r_sub.font.color.rgb = RGBColor(107, 114, 128)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("InternSphere\nAI-Powered Internship Matching & Career Assistant\n\n")
    r_title.font.name = "Calibri"
    r_title.font.bold = True
    r_title.font.size = Pt(26)
    r_title.font.color.rgb = RGBColor(17, 24, 39)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_runs = [
        ("Project Title:\n", True, 12),
        ("InternSphere — AI-Powered Internship Matching & Career Assistant\n\n", False, 12),
        ("Prepared By:\n", True, 12),
        ("[Intern Name]\n\n", False, 12),
        ("Project Duration:\n", True, 12),
        ("[Start Date – End Date]\n\n", False, 12),
        ("Mentor:\n", True, 12),
        ("[Mentor Name]\n\n", False, 12),
        ("Team Members:\n", True, 12),
        ("[Team Members]\n\n", False, 12),
        ("Submission Date:\n", True, 12),
        ("[Submission Date]\n", False, 12),
    ]
    for text, bold, size in meta_runs:
        r = p_meta.add_run(text)
        r.font.name = "Calibri"
        r.font.bold = bold
        r.font.size = Pt(size)
        r.font.color.rgb = RGBColor(31, 41, 55) if not bold else RGBColor(55, 65, 81)

    doc.add_page_break()

    # Page 2: Index
    h_idx = doc.add_heading("Index", level=1)
    h_idx.paragraph_format.space_after = Pt(12)

    index_headers = ["No.", "Section", "Page No."]
    index_rows = [
        ["1", "Project Overview", "3"],
        ["2", "Project Objective", "4"],
        ["3", "Technology Stack", "5"],
        ["4", "Project Requirements", "6"],
        ["5", "System Architecture / Workflow", "7"],
        ["6", "Product Backlog", "8"],
        ["7", "Sprint Planning", "9"],
        ["8", "Sprint Tasks & Status", "10"],
        ["9", "Implementation / Key Features", "11"],
        ["10", "Testing", "12"],
        ["11", "Challenges & Solutions", "13"],
        ["12", "Final Project Status", "14"],
        ["13", "10-Slide Project PPT Structure", "15"],
    ]
    tbl_index = doc.add_table(rows=1, cols=3)
    style_table(tbl_index, [0.8, 4.5, 1.2], index_headers, index_rows)

    doc.add_page_break()

    # 1. Project Overview
    doc.add_heading("1. Project Overview", level=1)
    
    doc.add_heading("1.1 Project Name", level=2)
    p = doc.add_paragraph("InternSphere — AI-Powered Internship Matching & Career Assistant")

    doc.add_heading("1.2 What is the Project?", level=2)
    p = doc.add_paragraph(
        "InternSphere is a full-stack, AI-augmented career acceleration and internship matching platform designed to streamline every phase of an intern's application lifecycle. The platform ingests candidate resumes (.pdf and .docx), runs a hybrid deterministic regex and Google Gemini extraction pipeline to build a structured profile, performs explainable semantic and skill-based internship matching over a curated vector index, conducts granular skill-gap analyses, drafts hallucination-free cover letters downloadable as PDFs, answers platform questions through a grounded AI Assistant, and provides an interactive, multi-turn AI Interview Preparation Agent equipped with role recommendations, roadmap synthesis, and browser-based voice read-aloud playback."
    )

    doc.add_heading("1.3 Problem Statement", level=2)
    p = doc.add_paragraph(
        "Internship seekers, college students, and entry-level graduates face substantial challenges when seeking practical industry experience:\n"
        "1. Inefficient Internship Discovery: Generic search filters fail to grasp candidate projects, coursework, and technical competencies, leading to poor role alignment.\n"
        "2. Opaque Skill Deficiencies: Candidates rarely receive actionable feedback regarding why they were rejected or what skills they lack for specific listings.\n"
        "3. Time-Consuming Application Customization: Tailoring unique cover letters and responses for multiple employers consumes significant effort and often leads to generic submissions.\n"
        "4. Disjointed Interview Preparation: Traditional interview study relies on static question lists without contextual adaptation to a candidate's actual projects or technical background."
    )

    doc.add_heading("1.4 Target Users", level=2)
    p = doc.add_paragraph(
        "• Students: Undergraduate and graduate students actively seeking summer or semester internships.\n"
        "• Internship Seekers & Fresh Graduates: Entry-level professionals looking to transition academic skills into corporate roles.\n"
        "• Career Changers: Candidates who require personalized skill-gap identification and targeted learning roadmaps to break into tech domains."
    )

    # 2. Project Objective
    doc.add_heading("2. Project Objective", level=1)
    p = doc.add_paragraph(
        "The primary objective of InternSphere is to provide an end-to-end, intelligent career companion that bridges the gap between candidate qualifications and industry internship openings.\n\n"
        "Objectives:\n"
        "• Automate Resume Parsing & Information Extraction: Eliminate manual data entry by extracting contact details, education, technical skills, soft skills, and projects from PDF and Word resumes using hybrid regex and LLM parsing.\n"
        "• Deliver Explainable Semantic Internship Matching: Rank internship postings against candidate qualifications using a hybrid scoring algorithm combining 384-dimensional vector similarity and explicit required/preferred skill overlap ratios.\n"
        "• Provide Clear Skill-Gap Visibility: Highlight missing required skills, preferred qualifications, and eligibility prerequisites directly on internship recommendation cards.\n"
        "• Automate Tailored Cover Letter Synthesis: Generate personalized, factual cover letters based strictly on candidate resume data and job facts, complete with in-browser editing and instant binary PDF generation.\n"
        "• Deliver Domain-Separated AI Career Coaching: Provide a dedicated AI Assistant for platform operations and an AI Interview Preparation Agent for mock interviews, STAR-method guidance, and voice playback."
    )

    # 3. Technology Stack
    doc.add_heading("3. Technology Stack", level=1)
    p = doc.add_paragraph("InternSphere is engineered strictly with lightweight, verified technologies:")
    tech_headers = ["Category", "Technology Used", "Purpose / Justification"]
    tech_rows = [
        ["Programming Language", "Python 3.12+", "Core backend language providing native async performance and typing."],
        ["Backend Framework", "FastAPI (v0.115+)", "High-performance ASGI framework for RESTful APIs and static file routing."],
        ["Frontend UI", "Vanilla HTML5 & CSS3", "Custom responsive design system with dark/light themes and glassmorphism."],
        ["Frontend Logic", "Vanilla JavaScript (ES6+)", "Modular asynchronous client-side state management without heavy frameworks."],
        ["Voice / Audio API", "Web Speech API", "Native browser window.speechSynthesis for read-aloud interview prep narration."],
        ["Database", "SQLite / PostgreSQL", "Local relational development on resume_parser.db; production-ready with PostgreSQL."],
        ["Database ORM", "SQLAlchemy 2.0+", "Declarative ORM models, relational cascades, and safe SQLite schema migrations."],
        ["Authentication", "JWT & bcrypt", "Stateless JSON Web Tokens (python-jose, HS256) and secure password hashing (passlib)."],
        ["AI / LLM Engine", "Google Gemini API", "Powered by gemini-3.5-flash-lite, gemini-3.6-flash via official google-genai SDK."],
        ["Vector Store / RAG", "Persisted JSON Index", "Lightweight 384-dimensional cosine similarity store (internship_vector_index.json)."],
        ["Embeddings", "Local Hashing / Gemini", "Fast SHA-256 word n-gram feature hashing (offline) or gemini-embedding-001."],
        ["Document Processing", "pdfplumber / python-docx", "PDF text extraction and DOCX document stream parsing."],
        ["PDF Generation", "Custom PDF-1.4 Generator", "Zero-dependency pure-Python binary stream generator for cover letters."],
        ["Testing", "Python unittest", "Automated test suite across candidate matching, API routes, and session memory."],
        ["Version Control", "Git & GitHub", "Distributed source code management and collaborative version control."],
    ]
    tbl_tech = doc.add_table(rows=1, cols=3)
    style_table(tbl_tech, [1.5, 2.0, 3.0], tech_headers, tech_rows)

    # 4. Project Requirements
    doc.add_heading("4. Project Requirements", level=1)
    doc.add_heading("Functional Requirements", level=2)
    p = doc.add_paragraph(
        "• User Authentication & Authorization: Users can register, log in, change passwords, and request password reset tokens; secured via JWT bearer tokens.\n"
        "• Profile Management: Users can update personal details, college, GPA, social links, technical skills, and upload profile pictures.\n"
        "• Resume Upload & Parsing: Users can upload .pdf or .docx resumes (up to 5 MB); system parses and stores structured candidate facts.\n"
        "• Semantic Internship Retrieval & Matching: System matches candidate resumes against curated internships and displays match scores with skill gap breakdowns.\n"
        "• Application Tracking: Candidates can submit, view, update, and withdraw applications with prevention of duplicate entries.\n"
        "• Cover Letter Generation & Export: System drafts role-tailored cover letters using candidate facts, allows in-browser edits, and exports to TXT and PDF.\n"
        "• Grounded AI Platform Assistant: Multi-turn chat assistant providing platform navigation and FAQ answers backed by documentation RAG.\n"
        "• AI Interview Preparation Coach: Generates Technical, HR, and Project questions, provides STAR answer advice, produces learning roadmaps, and speaks answers aloud.\n"
        "• Custom Study Document RAG: Users can upload domain-specific study notes or job descriptions for grounded interview prep."
    )
    doc.add_heading("Non-Functional Requirements", level=2)
    p = doc.add_paragraph(
        "• Security: Passwords encrypted with bcrypt; JWT access tokens expire in 60 minutes; owner-only access controls enforced across all personal data.\n"
        "• Performance & Latency: Sub-second deterministic local vector search; asynchronous background LLM invocations.\n"
        "• Reliability & Anti-Hallucination: Prompts enforce fact-bounded generation for cover letters and assistant answers.\n"
        "• Usability: Responsive single-page interface with real-time toast feedback, modal confirmation dialogs, and dark/light themes.\n"
        "• Maintainability: Modular architecture isolating database models, API routers, and business logic services."
    )

    # 5. System Architecture / Workflow
    doc.add_heading("5. System Architecture / Workflow", level=1)
    p = doc.add_paragraph(
        "InternSphere follows a clean multi-tier architecture uniting a Single Page Application frontend, FastAPI ASGI backend, SQLAlchemy ORM, and hybrid AI/RAG services:\n\n"
        "Architecture Flow:\n"
        "User Browser (HTML5 / CSS3 / JS / Web Speech API)\n"
        "  ↓  HTTP REST Requests (Bearer JWT)\n"
        "FastAPI Backend Application\n"
        "  ↓  Routers & Controllers\n"
        "Business Logic & Service Layer (Hybrid Parser / RAG Matcher / Cover Letter Generator / Interview Coach)\n"
        "  ↓  SQLAlchemy 2.0 ORM Engine\n"
        "Database Storage (SQLite: resume_parser.db / PostgreSQL) + AI Services (Google GenAI Gemini / Local Embeddings)\n"
        "  ↓  JSON Envelope / Binary PDF Stream\n"
        "User Browser"
    )
    p = doc.add_paragraph(
        "Key AI / RAG Pipelines:\n"
        "1. Resume Pipeline: Document Upload (.pdf/.docx) → Text Extraction (pdfplumber/docx) → Regex (Contacts) + Gemini LLM (Skills, Experience) → merge.py → resumes Table.\n"
        "2. Internship Matching Pipeline: Candidate Profile/Resume → 384-d Search Vector → Cosine Similarity Retrieval → Canonical Skill Ratio Evaluation → Ranked Output Cards.\n"
        "3. AI Assistant Pipeline: Question → Session Memory Context → RAG over product_knowledge.json → Gemini Prompt → Response + Sources Display.\n"
        "4. Interview Coaching Pipeline: Candidate Resume → Role Recommendations → Categorized Question Synthesis (Tech/HR/Project) → Answer Guidance & Roadmap → Web Speech Voice Playback."
    )

    # 6. Product Backlog
    doc.add_heading("6. Product Backlog", level=1)
    pb_headers = ["ID", "Feature / Task", "Module", "Priority", "Status"]
    pb_rows = [
        ["PB-01", "Project Initialization & FastAPI Static SPA Setup", "Core / Infra", "High", "Done"],
        ["PB-02", "Database Configuration (SQLAlchemy, SQLite, Safe Migrations)", "Database", "High", "Done"],
        ["PB-03", "User Authentication (Registration, Login, JWT Tokens, Password Reset)", "Auth", "High", "Done"],
        ["PB-04", "User Profile Management (Bio, Skills, Education, GPA, Picture Upload)", "Profile", "High", "Done"],
        ["PB-05", "Hybrid Resume Extraction Pipeline (pdfplumber, docx, Regex, Gemini)", "Resume", "High", "Done"],
        ["PB-06", "Vector Indexing & Ingestion Engine (internship_vector_index.json)", "RAG / Search", "High", "Done"],
        ["PB-07", "Explainable Internship Matching Algorithm & Skill-Gap Breakdown", "Matching", "High", "Done"],
        ["PB-08", "Application Lifecycle Management (Apply, Track, Update, Withdraw)", "Applications", "Medium", "Done"],
        ["PB-09", "AI Cover Letter Generator with Fact-Bounded Prompts & PDF Export", "Cover Letter", "Medium", "Done"],
        ["PB-10", "Grounded Product AI Assistant with Persistent Multi-Turn Chat Sessions", "AI Assistant", "Medium", "Done"],
        ["PB-11", "AI Interview Preparation Agent (Role Advice, Questions, Roadmaps)", "Interview Prep", "High", "Done"],
        ["PB-12", "Custom Study Document Upload & Retrieval for Interview Grounding", "Interview Prep", "Medium", "Done"],
        ["PB-13", "Web Speech API Audio Synthesis Integration (window.speechSynthesis)", "Frontend / UI", "Medium", "Done"],
        ["PB-14", "Comprehensive Automated Test Suite (unittest Discovery)", "Testing", "High", "Done"],
        ["PB-15", "Production Cloud Deployment & Containerization Setup", "DevOps", "Low", "To Do"],
    ]
    tbl_pb = doc.add_table(rows=1, cols=5)
    style_table(tbl_pb, [0.8, 2.5, 1.2, 1.0, 1.0], pb_headers, pb_rows)

    # 7. Sprint Planning
    doc.add_heading("7. Sprint Planning", level=1)
    sprints_info = [
        ("Sprint 1", "1 Week ([Start Date] – [Date + 7 Days])", "Foundation, Architecture & Authentication",
         ["Understand requirements & establish project architecture",
          "Configure SQLAlchemy with automatic table creation & SQLite foreign keys",
          "Implement /register, /login, /me, and password reset endpoints",
          "Create base Single Page Application layout with authentication modals"]),
        ("Sprint 2", "1 Week ([Date + 8 Days] – [Date + 14 Days])", "Resume Ingestion & Profile Management",
         ["Implement /profile CRUD endpoints and profile picture uploads",
          "Integrate pdfplumber and python-docx for document text extraction",
          "Build regex_extractor.py for deterministic contact and link parsing",
          "Connect llm_extractor.py for structured skill/experience extraction via Gemini",
          "Reconcile outputs via merge.py and persist into the resumes table"]),
        ("Sprint 3", "1 Week ([Date + 15 Days] – [Date + 21 Days])", "Semantic Internship Matching & Applications",
         ["Construct 384-dimensional feature hashing and Gemini embedding indexers",
          "Build hybrid matching algorithm combining cosine similarity (35%) and skill ratios (65%)",
          "Implement skill-gap calculations (matching, missing required, missing preferred)",
          "Create application lifecycle module (/applications) with duplicate prevention",
          "Design interactive matching cards and modal detail dialogs"]),
        ("Sprint 4", "1 Week ([Date + 22 Days] – [Date + 28 Days])", "AI Career Services, Cover Letters & Interview Coach",
         ["Build fact-bounded Gemini cover letter generation with custom PDF-1.4 downloads",
          "Develop Product AI Assistant using RAG over data/product_knowledge.json",
          "Implement Interview Preparation Agent with role recommendations and roadmaps",
          "Add database session persistence to retain multi-turn context",
          "Integrate browser Web Speech API for voice playback of coaching responses"]),
        ("Sprint 5", "1 Week ([Date + 29 Days] – [End Date])", "Testing, Hardening & Finalization",
         ["Execute test suites across candidate matching, session isolation, and interview APIs",
          "Implement safe schema migration utilities (apply_safe_schema_updates)",
          "Sanitize Gemini error responses to eliminate credential leak risks",
          "Complete technical documentation, Agile report, and PPT structure"]),
    ]
    for sp_name, sp_dur, sp_goal, sp_tasks in sprints_info:
        doc.add_heading(f"{sp_name} — {sp_goal}", level=2)
        p = doc.add_paragraph(f"Duration: {sp_dur}\nGoal: {sp_goal}\nTasks:")
        for t in sp_tasks:
            doc.add_paragraph(f"• {t}", style='List Bullet')

    # 8. Sprint Tasks & Status
    doc.add_heading("8. Sprint Tasks & Status", level=1)
    sp_headers = ["Sprint", "Task Description", "Assigned To", "Status"]
    sp_rows = [
        ["Sprint 1", "Project setup, folder conventions, and .env handling", "Intern / Developer", "Done"],
        ["Sprint 1", "SQLAlchemy engine & relational model definitions (models.py)", "Intern / Developer", "Done"],
        ["Sprint 1", "JWT token generation, bcrypt hashing, and auth endpoints", "Intern / Developer", "Done"],
        ["Sprint 1", "Base single-page frontend structure (index.html, styles.css)", "Intern / Developer", "Done"],
        ["Sprint 2", "Profile management endpoints & image upload pipeline", "Intern / Developer", "Done"],
        ["Sprint 2", "PDF & DOCX text extraction pipeline (extract_text.py)", "Intern / Developer", "Done"],
        ["Sprint 2", "Regex and Gemini LLM structured parsing integration", "Intern / Developer", "Done"],
        ["Sprint 2", "Profile & Resume management dashboard UI", "Intern / Developer", "Done"],
        ["Sprint 3", "Persisted JSON vector indexer (internship_index.py)", "Intern / Developer", "Done"],
        ["Sprint 3", "Hybrid scoring algorithm & skill-gap analysis logic", "Intern / Developer", "Done"],
        ["Sprint 3", "Application tracking APIs & status withdrawal endpoints", "Intern / Developer", "Done"],
        ["Sprint 3", "Interactive matching dashboard with skill tags and filters", "Intern / Developer", "Done"],
        ["Sprint 4", "Gemini cover letter generator & pure-Python PDF builder", "Intern / Developer", "Done"],
        ["Sprint 4", "Product AI Assistant with RAG knowledge retrieval", "Intern / Developer", "Done"],
        ["Sprint 4", "Interview Preparation Agent with tailored question generation", "Intern / Developer", "Done"],
        ["Sprint 4", "Multi-turn chat persistence in database tables", "Intern / Developer", "Done"],
        ["Sprint 4", "Web Speech API audio synthesis (🔊 Listen controls)", "Intern / Developer", "Done"],
        ["Sprint 5", "Automated test suite execution across all test files", "Intern / Developer", "Done"],
        ["Sprint 5", "Out-of-scope question routing & credential sanitization", "Intern / Developer", "Done"],
        ["Sprint 5", "Complete documentation, Agile report, and PPT structure", "Intern / Developer", "Done"],
    ]
    tbl_sp = doc.add_table(rows=1, cols=4)
    style_table(tbl_sp, [1.0, 3.2, 1.3, 1.0], sp_headers, sp_rows)

    # 9. Implementation / Key Features
    doc.add_heading("9. Implementation / Key Features", level=1)
    features = [
        ("Feature 1: User Authentication & Profile Management",
         "Secure account registration, authentication, password management, and personal profile administration.",
         "FastAPI routers (auth_routes.py, profile_routes.py) backed by passlib[bcrypt] and python-jose. User profiles support detailed education history, branch, graduation year, CGPA, and image uploads.",
         "[Insert Screenshot: Login and Profile Management Page]"),
        ("Feature 2: Hybrid Resume Parsing Pipeline",
         "Automatic parsing of candidate resumes in .pdf or .docx format into structured JSON.",
         "Text is extracted via pdfplumber or python-docx. Deterministic regular expressions parse contact details and URLs. Google Gemini extracts structured skill lists, employment history, and education. Results are unified and stored in SQLite.",
         "[Insert Screenshot: Resume Upload and Parsed Extraction View]"),
        ("Feature 3: Semantic Internship Matching & Skill-Gap Analysis",
         "Intelligent recommendation engine matching candidate resumes against internships, providing match percentages and skill gap breakdowns.",
         "Candidate facts are embedded into a 384-dimensional vector and compared against internship listings using cosine similarity. A weighted scoring function (35% semantic + 65% skill ratio) determines rankings. Missing required and preferred skills are clearly itemized.",
         "[Insert Screenshot: Internship Matching Cards and Skill-Gap Breakdown]"),
        ("Feature 4: Application Lifecycle Tracking",
         "Comprehensive application tracking allowing students to apply, review application statuses, and withdraw active submissions.",
         "The applications table links users to specific internships with attached resumes and cover letters. Unique constraints prevent duplicate submissions.",
         "[Insert Screenshot: Applications Tracker Dashboard]"),
        ("Feature 5: AI Cover Letter Generator & Binary PDF Export",
         "Generates tailored, factual cover letters for specific internships with inline editing and instant PDF downloads.",
         "Prompt engineering instructs Gemini to use strictly candidate and job facts without hallucinating qualifications. The letter is rendered to a clean standard PDF-1.4 file via services/pdf_generator.py.",
         "[Insert Screenshot: Cover Letter Editor and Download View]"),
        ("Feature 6: Grounded Product AI Assistant",
         "Real-time interactive assistant answering user queries regarding platform navigation, matching criteria, and application rules.",
         "Implements RAG over data/product_knowledge.json with multi-turn session persistence in chat_sessions and chat_messages.",
         "[Insert Screenshot: AI Assistant Chat Interface]"),
        ("Feature 7: AI Interview Preparation Coach with Voice Synthesis",
         "Resume-tailored interview coach delivering role recommendations, technical/behavioral questions, STAR answer advice, and roadmaps.",
         "Leverages candidate resume entities to formulate structured questions. Implements conversational memory, out-of-scope filters, and triggers the browser's Web Speech API (window.speechSynthesis) for read-aloud audio narration.",
         "[Insert Screenshot: Interview Prep Dashboard and Audio Playback Controls]"),
    ]
    for f_title, f_desc, f_impl, f_out in features:
        doc.add_heading(f_title, level=2)
        p = doc.add_paragraph()
        r = p.add_run("Description:\n")
        r.bold = True
        p.add_run(f"{f_desc}\n\n")
        r = p.add_run("Implementation:\n")
        r.bold = True
        p.add_run(f"{f_impl}\n\n")
        r = p.add_run("Output / Result Placeholder:\n")
        r.bold = True
        p.add_run(f"{f_out}\n")

    # 10. Testing
    doc.add_heading("10. Testing", level=1)
    p = doc.add_paragraph(
        "Testing was conducted using Python's standard library unittest framework to validate business logic, RAG retrieval quality, scoring formulas, and API endpoint security."
    )
    t_headers = ["Test ID", "Test Case Description", "Expected Result", "Actual Result", "Status"]
    t_rows = [
        ["T001", "Register user with valid email & password", "User record created, status 201", "User created in DB", "Pass"],
        ["T002", "Register duplicate username or email", "HTTP 409 Conflict returned", "HTTP 409 with error envelope", "Pass"],
        ["T003", "Login with valid credentials", "Status 200, JWT access token returned", "JWT token returned", "Pass"],
        ["T004", "Login with incorrect password", "HTTP 401 Unauthorized returned", "HTTP 401 returned", "Pass"],
        ["T005", "Update profile education and skills", "Status 200, updated JSON returned", "Profile records updated", "Pass"],
        ["T006", "Upload valid PDF resume", "File saved, text parsed, DB updated", "Resume parsed & saved", "Pass"],
        ["T007", "Upload invalid file type (.txt/.exe)", "HTTP 415 Unsupported Media Type", "HTTP 415 returned", "Pass"],
        ["T008", "Ingest internship dataset", "Vector embeddings generated & saved", "Index saved to JSON", "Pass"],
        ["T009", "Match candidate against backend roles", "Python/FastAPI roles ranked top", "Roles correctly prioritized", "Pass"],
        ["T010", "Verify skill-gap missing skill output", "Missing required skills listed", "Accurately displayed", "Pass"],
        ["T011", "Submit duplicate internship application", "HTTP 409 Conflict returned", "HTTP 409 returned", "Pass"],
        ["T012", "Withdraw pending application", "Status set to Withdrawn, timestamp recorded", "Withdrawn successfully", "Pass"],
        ["T013", "Generate cover letter without credentials", "Safe error message without API key leak", "Credential redacted", "Pass"],
        ["T014", "Download cover letter as PDF", "Valid PDF-1.4 binary stream returned", "Valid PDF rendered", "Pass"],
        ["T015", "Product Assistant question answering", "Grounded answer with source references", "Answer returned", "Pass"],
        ["T016", "Interview Prep role recommendations", "Relevant roles returned for candidate", "Accurate roles returned", "Pass"],
        ["T017", "Interview Prep out-of-scope query", "Politely declines non-career trivia", "Politely declined", "Pass"],
        ["T018", "Multi-turn interview conversation isolation", "Messages isolated to owning session", "Sessions isolated", "Pass"],
    ]
    tbl_t = doc.add_table(rows=1, cols=5)
    style_table(tbl_t, [0.8, 2.2, 1.8, 1.2, 0.6], t_headers, t_rows)

    p = doc.add_paragraph(
        "\nTesting Types Conducted:\n"
        "• Functional Testing: End-to-end user workflows (registration → upload → match → apply).\n"
        "• API & Schema Testing: HTTP status codes, headers, and Pydantic validation across all endpoints.\n"
        "• Integration Testing: Database persistence, cascade deletions, and relational foreign keys.\n"
        "• AI & RAG Verification: Anti-hallucination verification, deterministic feature hashing, and semantic ranking stability."
    )

    # 11. Challenges & Solutions
    doc.add_heading("11. Challenges & Solutions", level=1)
    ch_headers = ["Challenge Encountered", "Technical Solution Implemented"]
    ch_rows = [
        ["Inconsistent Resume Formats (PDF/DOCX)", "Built a dual-engine extractor using pdfplumber and python-docx coupled with deterministic regex for phone/emails and Gemini for nuanced skills."],
        ["Hallucination in Cover Letters", "Designed strict zero-shot system prompts instructing Gemini to reference only explicit candidate facts and job requirements."],
        ["Offline Vector Search Without External Services", "Developed a deterministic 384-dimensional feature hashing algorithm using SHA-256 and word n-grams, enabling vector search without external vector DBs."],
        ["Domain Separation Between AI Engines", "Implemented regex pattern matching in interview_prep_service.py to intercept platform questions and route users to the Product Assistant."],
        ["Audio Playback Without Heavy Dependencies", "Utilized the browser-native Web Speech API (window.speechSynthesis), avoiding heavy audio transcription libraries."],
        ["Safe Database Migrations in SQLite", "Implemented apply_safe_schema_updates() to dynamically check table columns and append nullable fields without dropping user data."],
        ["Preventing Duplicate Applications", "Enforced composite unique constraints (uq_application_user_internship) at the database level with clean HTTP 409 exception handling."],
    ]
    tbl_ch = doc.add_table(rows=1, cols=2)
    style_table(tbl_ch, [2.5, 4.0], ch_headers, ch_rows)

    # 12. Final Project Status
    doc.add_heading("12. Final Project Status", level=1)
    p = doc.add_paragraph(
        "Completed Features (✅):\n"
        "• ✅ User Registration, Login, JWT Authentication, and Password Management\n"
        "• ✅ User Profile Administration with Social Links, Academic Details, and Photo Uploads\n"
        "• ✅ Authenticated PDF and DOCX Resume Parsing with Hybrid Regex + Gemini Extraction\n"
        "• ✅ Persisted 384-Dimensional Vector Indexing with Offline Feature Hashing Support\n"
        "• ✅ Explainable Semantic Internship Matching with Hybrid Scoring Formulas\n"
        "• ✅ Granular Skill-Gap Analysis Highlighting Matching, Required, and Preferred Skills\n"
        "• ✅ End-to-End Application Lifecycle Management with Status Tracking and Withdrawal\n"
        "• ✅ AI-Powered Cover Letter Generator with In-Browser Editing and Binary PDF Export\n"
        "• ✅ Grounded Product AI Assistant with Multi-Turn Session Memory and Documentation RAG\n"
        "• ✅ AI Interview Preparation Coach with Role Advice, Questions, and STAR Evaluation\n"
        "• ✅ Custom Study Guide Upload and Grounded Q&A for Interview Preparation\n"
        "• ✅ Hands-Free Voice Synthesis via the Browser Web Speech API (🔊 Listen)\n"
        "• ✅ Automated Test Suite Covering Unit and Integration Scenarios\n\n"
        "In Progress (🔄):\n"
        "• 🔄 Advanced ATS Resume Optimization Scoring and Keyword Density Suggestions\n"
        "• 🔄 Export of Interview Prep Roadmaps into Calendar-Friendly Schedules\n\n"
        "Pending / Future Scope (📋):\n"
        "• 📋 Integration with External Job Boards (LinkedIn, Indeed APIs) for Real-Time Openings\n"
        "• 📋 WebRTC-Based Interactive Mock Video and Audio Interviews with Speech-to-Text (Whisper)\n"
        "• 📋 Recruiter Dashboard for Posting Listings and Managing Candidate Submissions\n\n"
        "Final Outcome Summary:\n"
        "The InternSphere platform has been successfully developed, integrated, and verified as a fully operational, end-to-end career acceleration solution. The application runs reliably on local environments using SQLite and is architected for instant deployment to PostgreSQL. All primary modules—including authentication, resume parsing, semantic matching, application tracking, cover letter drafting, product assistance, and interview preparation—function as specified."
    )

    # 13. 10-Slide Project PPT Structure
    doc.add_heading("13. 10-Slide Project PPT Structure", level=1)
    slides = [
        ("Slide 1: Project Title",
         "• Title: InternSphere — AI Career Companion Agent for Internship Matching and Interview Preparation\n"
         "• Subtitle: An Intelligent Full-Stack Career Acceleration Platform\n"
         "• Presenter: [Intern Name]\n"
         "• Track: Full-Stack & AI Engineering Intern\n"
         "• Mentor: [Mentor Name]\n"
         "• Date: [Submission Date]"),
        ("Slide 2: Problem Statement",
         "• Inefficient Job Discovery: Keyword search engines fail to understand candidate project context and coursework.\n"
         "• Opaque Skill Gaps: Applicants receive rejection emails without insights into required missing competencies.\n"
         "• Application Burnout: Crafting unique cover letters for dozens of applications is tedious and repetitive.\n"
         "• Generic Interview Prep: Traditional interview guides do not adapt to a candidate's actual projects or target roles."),
        ("Slide 3: Project Objectives",
         "• Build an end-to-end platform connecting student qualifications to tailored internship openings.\n"
         "• Automate hybrid resume parsing using deterministic regex and Google Gemini structured extraction.\n"
         "• Provide transparent, explainable matching via 384-dimensional semantic search and skill ratio calculations.\n"
         "• Synthesize hallucination-free cover letters downloadable as clean, standard PDF documents.\n"
         "• Provide tailored AI interview coaching with role recommendations, answer guidance, and voice playback."),
        ("Slide 4: Proposed Solution",
         "• Single-Page Web Application: Responsive, lightweight dashboard built with HTML5, CSS3, and JavaScript.\n"
         "• FastAPI Backend: Asynchronous RESTful API engine orchestrating authentication, business logic, and RAG.\n"
         "• Hybrid Extraction: Two-tier parser combining regex precision with Gemini semantic comprehension.\n"
         "• Explainable Vector Matching: 384-d feature hashing index paired with canonical skill overlap scoring.\n"
         "• Dual AI Assistants: Dedicated Product Assistant for platform help and Interview Coach for career preparation."),
        ("Slide 5: Technology Stack",
         "• Backend: Python 3.12+, FastAPI, Uvicorn, SQLAlchemy 2.0.\n"
         "• Frontend: Vanilla HTML5, CSS3, ES6+ JavaScript, Web Speech API.\n"
         "• Database: SQLite (Development) / PostgreSQL (Production).\n"
         "• AI & LLM: Google GenAI SDK (gemini-3.5-flash-lite, gemini-3.6-flash).\n"
         "• Document Processing: pdfplumber, python-docx, Custom PDF-1.4 Generator.\n"
         "• Security: JWT Bearer Tokens (python-jose), bcrypt Password Hashing."),
        ("Slide 6: Key Features & Capabilities",
         "• Resume Ingestion: Instant parsing of PDF and Word resumes into structured candidate records.\n"
         "• Grounded Matching: Automated recommendation rankings showing percentage match and skill breakdowns.\n"
         "• Skill Gap Analysis: Clear visualization of matched skills, missing required skills, and preferred skills.\n"
         "• Instant Cover Letters: One-click generation based on verified facts with inline editing and PDF export.\n"
         "• Interactive Interview Coach: Role advice, categorized questions (Tech/HR/Project), and voice read-aloud."),
        ("Slide 7: Agile Methodology & Sprints",
         "• Agile Scrum Framework: 5 focused weekly sprints from initialization to hardening.\n"
         "• Sprint 1: Setup, architecture, database models, and JWT authentication.\n"
         "• Sprint 2: Profile management and hybrid PDF/DOCX resume extraction pipeline.\n"
         "• Sprint 3: Vector indexing, semantic internship matching, and application tracking.\n"
         "• Sprint 4: Cover letters, Product AI Assistant, and Interview Preparation Agent.\n"
         "• Sprint 5: Comprehensive test suites, security hardening, and final documentation."),
        ("Slide 8: System Architecture & Workflow",
         "• Client Layer: Single Page App communicating via secure JWT Bearer REST headers.\n"
         "• Application Layer: FastAPI routes delegating to specialized services and controllers.\n"
         "• Persistence Layer: SQLAlchemy ORM managing 10 relational tables with cascade isolation.\n"
         "• AI / RAG Pipeline: Local 384-d vector embeddings, cosine retrieval, and Gemini prompt synthesis.\n"
         "• Audio Layer: Browser Web Speech API providing hands-free coaching playback."),
        ("Slide 9: Testing & Results",
         "• Test Suite: Automated unittest test suite covering matching, API routes, and session memory.\n"
         "• Functional Verification: Validated full candidate journey from registration to interview preparation.\n"
         "• Security Hardening: Sanitized error envelopes to prevent credential or API key leakage.\n"
         "• Reliability: Verified fact-bounded cover letter synthesis preventing hallucinated qualifications."),
        ("Slide 10: Conclusion & Future Scope",
         "• Project Success: Delivered a robust, self-contained AI platform resolving core internship search hurdles.\n"
         "• Key Takeaways: Practical mastery of FastAPI, SQLAlchemy ORM, RAG vector indexing, and Gemini APIs.\n"
         "• Future Scope: Live job board scraping (LinkedIn API), speech-to-text mock video interviews, and ATS resume scoring."),
    ]
    for s_title, s_body in slides:
        doc.add_heading(s_title, level=2)
        doc.add_paragraph(s_body)

    out_path = "Agile_Project_Documentation.docx"
    doc.save(out_path)
    print(f"Successfully generated {out_path}")


if __name__ == "__main__":
    build_document()
